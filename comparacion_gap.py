# -*- coding: utf-8 -*-
"""
Fase 3: comparación de gap genérico (TRFPO, PEAPO, GUIAPO, o cualquier
otro que se agregue en el futuro) vs cada código específico, a partir
de la salida de Fase 1 (extraccion_tarifas_vigentes.py,
tarifas_vigentes.xlsx).

A diferencia del diseño original en BRIEF.md (pensado antes de tener
rangos de pax reales leídos de Tourplan), Fase 1 ya trae el
PAX_DESDE/PAX_HASTA real de cada código directamente de la grilla de
RATES. Pero el PAX_DESDE/HASTA que carga el ESPECÍFICO en su propia
tarifa no necesariamente coincide con el tramo de pax que le
corresponde a su vehículo (bug real encontrado por la usuaria: un
código con sufijo "19" (Sprinter 19 pax) cargado como "1-11 pax" se
comparaba contra el tramo de TRFPO de Auto por simple superposición
numérica, en vez de contra el tramo de Sprinter 19 — ver DISENO.md "Bug
real: matching por vehículo, no sólo por pax bruto"). El cruce
completo:

1. Código específico → código genérico: longest-prefix-match, reusando
   tal cual mejor_prefijo_trfpo de matching_engine.py (Fase 2) — el
   algoritmo es puro string matching y nunca dependió de que el
   genérico fuera TRFPO en particular.
2. Código específico → vehículo: por el SUFIJO del código (ver
   SUFIJO_VEHICULO — los sufijos están estandarizados, mismo
   significado siempre, DISENO.md "Corrección a BRIEF.md"), no por
   Description. V9 (Van 9 pax) es un vehículo real que TRFPO no
   valoriza como tramo propio — se excluye de la comparación.
3. Código específico → categoría (EXCURSION/TRASLADO/JAPON/CRUCEROS_*)
   y guía (CON/SIN_GUIA): por Description, reusando
   detectar_flags/clasificar_categoria/clasificar_guia de
   matching_engine.py (Fase 2, sin cambios).
4. Vehículo + categoría + guía → rango de pax REAL de ese vehículo,
   vía tabla_bases_vehiculo_pax.csv (cargar_tabla_bases, también de
   matching_engine.py) — este es el rango que se usa para elegir el
   tramo de TRFPO, no el PAX_DESDE/HASTA que cargó el específico.
5. Código genérico + ese match → fila del genérico cuyo rango de pax
   REAL del vehículo se superpone (máxima superposición de pax; ante
   empate, máxima superposición de fechas con el período del
   específico) con alguno de sus tramos. Si `tarifas_vigentes.xlsx`
   trae la columna `GENERICO` (Fase 1 la agrega desde que se
   generalizó a PEAPO/GUIAPO), el cruce además se restringe al
   genérico de esa relación puntual — así, si dos genéricos comparten
   LOCATION (ej. TRFPO y PEAPO ambos en BUE), no se mezclan sus filas.
   Si `tarifas_vigentes.xlsx` no trae `DESCRIPCION` (Excel viejo, de
   antes de este fix) o el código tiene un sufijo no contemplado en
   SUFIJO_VEHICULO, se cae al comportamiento anterior: comparar
   directo contra el PAX_DESDE/HASTA que reportó el específico.
6. Gap = TARIFA_USD del específico vs. TARIFA_USD del genérico.

NOTA: `extraccion_tarifas_vigentes.py` (Fase 1) ya corre esta misma
comparación automáticamente al final de su propia ejecución (versión
en Python puro, sin pandas, embebida ahí para que una sola corrida de
Colab genere los dos Excel de una — ver DISENO.md, incluye su propia
copia de SUFIJO_VEHICULO/tabla_bases_vehiculo_pax.csv/clasificación por
Description porque corre como celda única sin archivos hermanos). Este
archivo queda para volver a correr SÓLO la comparación sin re-scrapear
Tourplan (ej. después de editar `tarifas_vigentes.xlsx` a mano, o para
reprocesar con otro archivo de entrada).

Requiere pandas + openpyxl. Uso:
    python comparacion_gap.py [tarifas_vigentes.xlsx] [salida.xlsx]
(default: tarifas_vigentes.xlsx → comparacion_gap.xlsx)
"""
import os
import sys
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

# El motor de matching (Fase 2) sigue llamándose mejor_prefijo_trfpo
# porque matching_engine.py es intencionalmente específico del dominio
# de transporte (ver DISENO.md) y no se generalizó — pero el algoritmo
# en sí es puro string-prefix-match, así que se reusa tal cual para
# cualquier relación genérico/específico. detectar_flags/
# clasificar_categoria/clasificar_guia/cargar_tabla_bases también se
# reusan tal cual (Fase 2, sin cambios) para clasificar cada código
# específico y saber el rango de pax real de su vehículo.
from matching_engine import (
    mejor_prefijo_trfpo as mejor_prefijo_generico,
    detectar_flags, clasificar_categoria, clasificar_guia,
    cargar_tabla_bases,
)

COLS_GAP = [
    "LOCATION", "CODIGO_GENERICO", "SUPPLIER_ESPECIFICO", "CODIGO_ESPECIFICO",
    "PAX_DESDE", "PAX_HASTA",
    "PERIODO_ESPECIFICO_DESDE", "PERIODO_ESPECIFICO_HASTA",
    "PERIODO_GENERICO_DESDE", "PERIODO_GENERICO_HASTA",
    "TARIFA_GENERICO_USD", "TARIFA_ESPECIFICO_USD",
    "DIFERENCIA_USD", "DIFERENCIA_PCT", "FLAGS",
]

TABLA_BASES_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "config", "tabla_bases_vehiculo_pax.csv")

# Sufijo de código -> nombre de vehículo, EXACTO como aparece en la
# columna VEHICULO de tabla_bases_vehiculo_pax.csv. Ojo: no usa
# normalizar_vehiculo de matching_engine.py acá — ese normalizador
# colapsa cualquier cosa que empiece con "bus" (incluida "Bus 37
# asientos") a la clave genérica "BUS", lo que perdería la distinción
# entre Bus 37 asientos (sufijo "37") y Bus/42 asientos (sufijo "42")
# que la tabla sí hace — por eso acá se filtra por el texto crudo de
# VEHICULO, no por VEHICULO_NORM. V9 (Van 9 pax) es un vehículo real
# que TRFPO no valoriza como tramo propio (confirmado por la usuaria)
# — se excluye de la comparación en vez de forzar un match.
SUFIJO_VEHICULO = {
    "AU": "Auto",
    "MV": "H1/VITO",
    "15": "Sprinter 15 pax",
    "19": "Sprinter 19 pax",
    "24": "Minibus",
    "37": "Bus 37 asientos",
    "42": "Bus",
}
SUFIJOS_VEHICULO_NO_VALORIZADOS_TRFPO = {"V9"}


def _vehiculo_por_sufijo(codigo):
    """Devuelve (vehiculo, sufijo) — ver el docstring de SUFIJO_VEHICULO."""
    sufijo = str(codigo or "")[-2:].upper()
    if sufijo in SUFIJOS_VEHICULO_NO_VALORIZADOS_TRFPO:
        return "NO_VALORIZADO", sufijo
    return SUFIJO_VEHICULO.get(sufijo), sufijo


def _cargar_tabla_bases_o_none():
    """None si el archivo no está disponible (ej. corriendo este
    script fuera del repo, sin config/) — construir_comparacion cae al
    comportamiento anterior (pax bruto del específico) en ese caso, en
    vez de romper toda la corrida."""
    try:
        return cargar_tabla_bases(TABLA_BASES_PATH)
    except (FileNotFoundError, OSError):
        print(f"    ⚠ No encontré {TABLA_BASES_PATH} — el matching por vehículo "
              f"queda deshabilitado, se compara por el pax bruto que reportó cada "
              f"específico (comportamiento anterior).")
        return None


def _pax_esperado_vehiculo(df_bases, location, categoria, vehiculo, guia):
    """Busca en df_bases (cargar_tabla_bases) el rango de pax real de
    un vehículo para esa location/categoría/guía — filtra por VEHICULO
    crudo, no por VEHICULO_NORM (ver SUFIJO_VEHICULO). Si no hay fila
    para esa guía puntual, cae a buscar sin filtrar por guía."""
    cand = df_bases[(df_bases["LOCATION"] == location) & (df_bases["CATEGORIA"] == categoria) &
                     (df_bases["VEHICULO"] == vehiculo) & (df_bases["GUIA"] == guia)]
    if cand.empty:
        cand = df_bases[(df_bases["LOCATION"] == location) & (df_bases["CATEGORIA"] == categoria) &
                         (df_bases["VEHICULO"] == vehiculo)]
    if cand.empty:
        return None
    r = cand.iloc[0]
    return int(r["PAX_DESDE"]), int(r["PAX_HASTA"])


def _pax_overlap(desde1, hasta1, desde2, hasta2):
    """Cantidad de pax en común entre dos rangos — análogo a
    _dias_superposicion pero para rangos de pax en vez de fechas."""
    if desde1 is None or hasta1 is None or pd.isna(desde2) or pd.isna(hasta2):
        return 0
    return max(0, min(hasta1, hasta2) - max(desde1, desde2) + 1)


def _parsear_fecha(txt):
    if not txt or (isinstance(txt, float) and pd.isna(txt)):
        return None
    if isinstance(txt, datetime):
        return txt
    try:
        return datetime.strptime(str(txt).strip(), "%d/%m/%Y")
    except ValueError:
        return None


def _r2(x):
    """Redondeo defensivo a 2 decimales — el Excel de entrada puede
    traer más decimales si viene de una división (ej. conversión ARS
    con un tipo de cambio no redondo)."""
    return round(x, 2) if pd.notna(x) else None


def _dias_superposicion(desde1, hasta1, desde2, hasta2):
    """Días de superposición entre dos rangos de fechas. 0 (no
    negativo) si no se superponen o si falta alguna fecha — para poder
    comparar candidatos con max() sin excepciones por None."""
    if not (desde1 and hasta1 and desde2 and hasta2):
        return 0
    inicio = max(desde1, desde2)
    fin = min(hasta1, hasta2)
    return max(0, (fin - inicio).days)


def cargar_tarifas(path):
    df = pd.read_excel(path)
    # Descarta el catch-all de pax abierto hacia arriba (ej. "42-9999
    # AD", siempre 0) — defensivo, por si el Excel de entrada viene de
    # una corrida vieja de extraccion_tarifas_vigentes.py que todavía
    # no filtraba esto en el origen.
    df = df[df["PAX_HASTA"] != 9999].reset_index(drop=True)
    df["_PERIODO_DESDE_DT"] = df["PERIODO_DESDE"].apply(_parsear_fecha)
    df["_PERIODO_HASTA_DT"] = df["PERIODO_HASTA"].apply(_parsear_fecha)
    return df


def construir_comparacion(df, df_bases=None):
    """df: DataFrame con las columnas de tarifas_vigentes.xlsx (una
    fila por código+período+rango de pax). df_bases: resultado de
    cargar_tabla_bases (o None para deshabilitar el matching por
    vehículo y comparar directo por el pax bruto del específico, como
    antes). Devuelve un DataFrame con una fila por (rango de pax,
    período) de CADA código específico, comparado contra el genérico
    que le corresponde."""
    filas_salida = []
    tiene_col_generico = "GENERICO" in df.columns
    tiene_col_descripcion = "DESCRIPCION" in df.columns

    for location, df_loc in df.groupby("LOCATION"):
        df_generico_loc = df_loc[df_loc["ES_GENERICO"] == True]
        df_especifico_loc = df_loc[df_loc["ES_GENERICO"] == False]

        for _, fila_e in df_especifico_loc.iterrows():
            codigo_e = str(fila_e["PRODUCT_CODE"])

            # Si el Excel trae GENERICO, se restringe el pool de
            # candidatos genéricos al de esta relación puntual — evita
            # mezclar TRFPO/PEAPO/GUIAPO cuando comparten LOCATION.
            # Si no lo trae (Excel de antes de la generalización), se
            # cae al comportamiento anterior de agrupar sólo por
            # LOCATION.
            if tiene_col_generico and pd.notna(fila_e.get("GENERICO")):
                df_generico = df_generico_loc[df_generico_loc["GENERICO"] == fila_e["GENERICO"]]
                if df_generico.empty:
                    df_generico = df_generico_loc
            else:
                df_generico = df_generico_loc

            codigos_generico = sorted(df_generico["PRODUCT_CODE"].astype(str).unique().tolist(),
                                       key=len, reverse=True)

            base = {
                "LOCATION": location,
                "SUPPLIER_ESPECIFICO": fila_e["SUPPLIER"],
                "CODIGO_ESPECIFICO": codigo_e,
                "PAX_DESDE": fila_e["PAX_DESDE"],
                "PAX_HASTA": fila_e["PAX_HASTA"],
                "PERIODO_ESPECIFICO_DESDE": fila_e["PERIODO_DESDE"],
                "PERIODO_ESPECIFICO_HASTA": fila_e["PERIODO_HASTA"],
                "PERIODO_GENERICO_DESDE": "", "PERIODO_GENERICO_HASTA": "",
                "TARIFA_ESPECIFICO_USD": _r2(fila_e["TARIFA_USD"]),
            }
            flags = []

            mejor, candidatos = mejor_prefijo_generico(codigo_e, codigos_generico)
            if len(candidatos) > 1:
                flags.append("COLISION_REVISAR")

            vehiculo, sufijo = _vehiculo_por_sufijo(codigo_e)
            if vehiculo == "NO_VALORIZADO":
                # Vehículo real (ej. V9/Van 9 pax) que TRFPO no
                # valoriza como tramo propio — no hay ningún tramo
                # contra el cual comparar, se excluye en vez de forzar
                # un match contra un tramo que no le corresponde.
                flags.append(f"VEHICULO_{sufijo}_NO_VALORIZADO_TRFPO")
                filas_salida.append({**base, "CODIGO_GENERICO": mejor or "",
                                      "TARIFA_GENERICO_USD": None,
                                      "DIFERENCIA_USD": None, "DIFERENCIA_PCT": None,
                                      "FLAGS": ",".join(flags)})
                continue

            if mejor is None:
                filas_salida.append({**base, "CODIGO_GENERICO": "",
                                      "TARIFA_GENERICO_USD": None,
                                      "DIFERENCIA_USD": None, "DIFERENCIA_PCT": None,
                                      "FLAGS": ",".join(flags + ["SIN_MATCH_GENERICO"])})
                continue

            # Rango de pax a usar para elegir el tramo correcto de
            # TRFPO: el "real" del vehículo según tabla_bases_vehiculo_
            # pax.csv (por sufijo de código + categoría/guía
            # clasificadas de la Description), no el PAX_DESDE/HASTA
            # que reportó el específico en su propia tarifa — puede ser
            # un tramo más ancho o distinto del que realmente le
            # corresponde a ese vehículo (ver DISENO.md).
            pax_desde_cmp, pax_hasta_cmp = fila_e["PAX_DESDE"], fila_e["PAX_HASTA"]
            if df_bases is None:
                flags.append("MATCHING_POR_VEHICULO_DESHABILITADO")
            elif not tiene_col_descripcion:
                flags.append("SIN_DESCRIPCION_EN_EXCEL")
            elif vehiculo is None:
                flags.append(f"SUFIJO_VEHICULO_DESCONOCIDO_{sufijo}")
            else:
                descripcion = fila_e.get("DESCRIPCION", "") or ""
                flags_codigo = detectar_flags(codigo_e, descripcion)
                categoria = clasificar_categoria(descripcion, flags_codigo)
                guia = clasificar_guia(flags_codigo, categoria)
                esperado = _pax_esperado_vehiculo(df_bases, location, categoria, vehiculo, guia)
                if esperado is None:
                    flags.append(f"SIN_BASE_VEHICULO_PAX_{categoria}_{vehiculo}".replace(" ", "_"))
                else:
                    pax_desde_cmp, pax_hasta_cmp = esperado

            candidatas_generico = df_generico[
                (df_generico["PRODUCT_CODE"].astype(str) == mejor) &
                (df_generico["PAX_DESDE"] <= pax_hasta_cmp) &
                (df_generico["PAX_HASTA"] >= pax_desde_cmp)
            ]
            if candidatas_generico.empty:
                filas_salida.append({**base, "CODIGO_GENERICO": mejor,
                                      "TARIFA_GENERICO_USD": None,
                                      "DIFERENCIA_USD": None, "DIFERENCIA_PCT": None,
                                      "FLAGS": ",".join(flags + ["SIN_GENERICO_PARA_ESE_PAX"])})
                continue

            # Más de un tramo de TRFPO candidato (ej. dos períodos
            # solapados, o el rango de pax esperado del vehículo cae a
            # caballo entre dos tramos): se elige primero por MAYOR
            # superposición de pax con el rango esperado del vehículo,
            # y ante empate por MAYOR superposición de fechas con el
            # período del específico. Se marca con
            # MULTIPLES_PERIODOS_GENERICO para no descartar el resto en
            # silencio.
            if len(candidatas_generico) > 1:
                flags.append("MULTIPLES_PERIODOS_GENERICO")

            fila_generico = max(
                candidatas_generico.to_dict("records"),
                key=lambda r: (
                    _pax_overlap(pax_desde_cmp, pax_hasta_cmp, r["PAX_DESDE"], r["PAX_HASTA"]),
                    _dias_superposicion(
                        fila_e["_PERIODO_DESDE_DT"], fila_e["_PERIODO_HASTA_DT"],
                        r["_PERIODO_DESDE_DT"], r["_PERIODO_HASTA_DT"])))

            if _dias_superposicion(fila_e["_PERIODO_DESDE_DT"], fila_e["_PERIODO_HASTA_DT"],
                                    fila_generico["_PERIODO_DESDE_DT"], fila_generico["_PERIODO_HASTA_DT"]) == 0:
                flags.append("SIN_SUPERPOSICION_DE_PERIODO")

            tarifa_generico = _r2(fila_generico["TARIFA_USD"])
            tarifa_especifico = _r2(fila_e["TARIFA_USD"])
            diff_usd = diff_pct = None
            if pd.notna(tarifa_generico) and pd.notna(tarifa_especifico):
                diff_usd = round(tarifa_especifico - tarifa_generico, 2)
                if tarifa_generico:
                    diff_pct = round((tarifa_especifico / tarifa_generico - 1) * 100, 2)
            else:
                flags.append("SIN_TARIFA_USD_PARA_COMPARAR")

            filas_salida.append({
                **base, "CODIGO_GENERICO": mejor,
                "PERIODO_GENERICO_DESDE": fila_generico["PERIODO_DESDE"],
                "PERIODO_GENERICO_HASTA": fila_generico["PERIODO_HASTA"],
                "TARIFA_GENERICO_USD": tarifa_generico,
                "DIFERENCIA_USD": diff_usd, "DIFERENCIA_PCT": diff_pct,
                "FLAGS": ",".join(flags),
            })

    return pd.DataFrame(filas_salida, columns=COLS_GAP)


def _formatear_columna_porcentaje(path_xlsx, nombre_columna):
    """Aplica formato de número con signo % (visual, sin multiplicar
    por 100 — DIFERENCIA_PCT ya está en escala porcentual, ej. -44.76
    significa -44.76%, no -0.4476) a una columna del Excel ya guardado.
    pandas.to_excel no permite formato por celda directamente, así que
    se reabre con openpyxl para aplicarlo y se vuelve a guardar."""
    wb = load_workbook(path_xlsx)
    hoja = wb.active
    headers = [c.value for c in hoja[1]]
    if nombre_columna not in headers:
        return
    idx = headers.index(nombre_columna) + 1
    for fila in hoja.iter_rows(min_row=2, min_col=idx, max_col=idx):
        for celda in fila:
            if isinstance(celda.value, (int, float)):
                celda.number_format = '0.00"%"'
    wb.save(path_xlsx)


def main():
    entrada = sys.argv[1] if len(sys.argv) > 1 else "tarifas_vigentes.xlsx"
    salida = sys.argv[2] if len(sys.argv) > 2 else "comparacion_gap.xlsx"
    df = cargar_tarifas(entrada)
    df_bases = _cargar_tabla_bases_o_none()
    comparacion = construir_comparacion(df, df_bases)
    comparacion.to_excel(salida, index=False, sheet_name="COMPARACION_GAP")
    _formatear_columna_porcentaje(salida, "DIFERENCIA_PCT")
    print(f"{len(comparacion)} filas escritas en {salida}")

    con_flags = comparacion[comparacion["FLAGS"] != ""]
    if not con_flags.empty:
        print(f"\n{len(con_flags)} filas con alguna bandera (revisar):")
        print(con_flags["FLAGS"].value_counts().to_string())


if __name__ == "__main__":
    main()
