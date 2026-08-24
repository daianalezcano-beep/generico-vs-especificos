# ============================================================
# GENÉRICO vs ESPECÍFICOS — FASE 1: extracción de tarifa vigente
# Google Colab — celda única (mismo formato que los scripts hermanos)
# ------------------------------------------------------------
#   Adaptado de:
#   - tourplan_valorizacion_pkg_v3.py (Tourplan-Valorizacion-EX-TF):
#     buscar_producto, _leer_tabla_rates, _extraer_valores_ad,
#     detección/instalación de Google Chrome en Colab (no viene
#     preinstalado), helpers base (ss/jc/set_val/wait/waitx/
#     crear_driver/login/hamburger/menu_item).
#   - cbd_impact_simulator_pkg.py (cbd-impact-simulator-pkg):
#     versión de buscar_producto que además abre el resultado y
#     confirma el contexto de producto (_en_contexto_producto),
#     STYPE_SIDEBAR.
#   Ver DISENO.md ("Reuso de las herramientas hermanas") para el
#   detalle de qué se tomó de cada una y qué se dejó afuera
#   (USED IN, PCM/Dashboard, subcode de moneda, recálculo, reversión —
#   nada de eso aplica: acá sólo se lee la tarifa vigente tal como está
#   cargada en el propio componente, sin simular nada).
#
#   GENERALIZADO A CUALQUIER PAR GENÉRICO/ESPECÍFICO: nació para TRFPO
#   (Transporte Por Asignar, service type TR) vs. transportistas, pero
#   COMPARACIONES acepta cualquier otra relación con la misma forma —
#   ej. PEAPO (service type PJ) o GUIAPO (service type GU) vs. sus
#   proveedores específicos — agregando otra entrada a la lista. Cada
#   relación puede tener su propio Price Code (ver "price_code" en
#   COMPARACIONES) — no hace falta que todas usen "TR". Ver DISENO.md.
#
#   SIN CSV/EXCEL DE ENTRADA: no hace falta exportar nada de Tourplan
#   antes de correr. Se edita COMPARACIONES (más abajo) con Location +
#   Supplier genérico + Supplier(es) específico(s), y el script busca
#   él mismo, en Tourplan, todos los códigos vigentes de cada supplier
#   (listar_codigos_supplier) y les lee la tarifa.
#
#   VALIDADO CONTRA TOURPLAN REAL (BUE, TRFPO + 6HOUS1): última corrida
#   completa trajo 116 códigos TRFPO + 52 de 6HOUS1, matching limpio en
#   204/208 filas de la comparación. Varios bugs reales ya encontrados
#   y corregidos en el camino — ver DISENO.md para el detalle de cada
#   uno. Sin confirmar todavía: el texto exacto del header de moneda en
#   la lista de períodos (`_leer_periodos_rates`); si conviene generar
#   una fila por cada período del genérico cuando hay más de uno
#   superpuesto (hoy se elige sólo el de mayor superposición, marcado
#   con MULTIPLES_PERIODOS_GENERICO); y todo lo específico de PJ/GU
#   (Price Code, exclusiones, formato de RATES) — sin corridas reales
#   contra esas relaciones todavía, sólo TRFPO está confirmado.
# ============================================================

import os, sys, subprocess, importlib, shutil, time, re
from datetime import datetime

print("🔧 Verificando entorno...\n")

# ── Paquetes Python ──────────────────────────────────────────
_PIPS_NEEDED = {
    "selenium":          "selenium",
    "webdriver_manager": "webdriver-manager",
    "openpyxl":          "openpyxl",
}
_pips_faltantes = [pkg for mod, pkg in _PIPS_NEEDED.items()
                   if importlib.util.find_spec(mod) is None]
if _pips_faltantes:
    subprocess.run([sys.executable, "-m", "pip", "install", "-q"] + _pips_faltantes, check=True)

# ── Google Chrome (REUTILIZADO tal cual de tourplan_valorizacion_pkg_v3.py
# — Colab NO trae Chrome preinstalado, hay que detectarlo/instalarlo cada
# corrida porque Colab resetea el entorno cuando pierde actividad) ──────

def _chrome_version(path):
    try:
        r = subprocess.run([path, "--version"], capture_output=True, text=True, timeout=8)
        v = (r.stdout or r.stderr).strip()
        return v if v and any(c.isdigit() for c in v) else ""
    except Exception:
        return ""


def _find_chrome():
    for name in ["google-chrome-stable", "google-chrome", "chromium-browser", "chromium"]:
        p = shutil.which(name)
        if p:
            v = _chrome_version(p)
            if v:
                return p, v
    for path in ["/usr/bin/google-chrome-stable", "/usr/bin/google-chrome"]:
        if os.path.exists(path):
            v = _chrome_version(path)
            if v:
                return path, v
    return None, ""


CHROMIUM_BIN, ver_chrome = _find_chrome()


def _instalar_chrome():
    """Instala google-chrome-stable en Colab. Estrategia 1: repo oficial
    de Google. Estrategia 2: descarga directa del .deb (fallback)."""
    DEVNULL = subprocess.DEVNULL

    def _apt_update():
        subprocess.run(["apt-get", "update", "-qq"], stdout=DEVNULL, stderr=DEVNULL)

    e1 = None
    try:
        print("  ⏳ Agregando repo Google Chrome y ejecutando apt...")
        subprocess.run(
            "wget -qO- https://dl.google.com/linux/linux_signing_key.pub "
            "| gpg --dearmor -o /usr/share/keyrings/google-chrome.gpg",
            shell=True, check=True, stdout=DEVNULL, stderr=DEVNULL)
        with open("/etc/apt/sources.list.d/google-chrome.list", "w") as f:
            f.write("deb [arch=amd64 signed-by=/usr/share/keyrings/google-chrome.gpg] "
                    "https://dl.google.com/linux/chrome/deb/ stable main\n")
        _apt_update()
        subprocess.run(["apt-get", "install", "-y", "-qq", "google-chrome-stable"],
                       check=True, stdout=DEVNULL, stderr=DEVNULL)
        if _find_chrome()[0]:
            return
    except Exception as e:
        e1 = e
        print(f"  ⚠ Repo Google falló ({e1}), probando descarga directa...")

    try:
        deb = "/tmp/google-chrome-stable.deb"
        url = "https://dl.google.com/linux/direct/google-chrome-stable_current_amd64.deb"
        print(f"  ⏳ Descargando {url} ...")
        subprocess.run(["wget", "-q", "-O", deb, url], check=True)
        print("  ⏳ Instalando .deb ...")
        subprocess.run(["dpkg", "-i", deb], stdout=DEVNULL, stderr=DEVNULL)
        _apt_update()
        subprocess.run(["apt-get", "install", "-f", "-y", "-qq"],
                       check=True, stdout=DEVNULL, stderr=DEVNULL)
        print("  ✅ Chrome instalado via .deb")
    except Exception as e2:
        raise EnvironmentError(
            f"No se pudo instalar google-chrome-stable.\n"
            f"  Repo: {e1 if e1 else 'n/a'}\n"
            f"  .deb:  {e2}\n"
            "Intentá manualmente en Colab:\n"
            "  !wget -q https://dl.google.com/linux/direct/google-chrome-stable_current_amd64.deb\n"
            "  !dpkg -i google-chrome-stable_current_amd64.deb\n"
            "  !apt-get install -f -y")


if not CHROMIUM_BIN:
    print("  ⏳ Chrome no encontrado, instalando...")
    _instalar_chrome()
    CHROMIUM_BIN, ver_chrome = _find_chrome()
    if CHROMIUM_BIN:
        print(f"  ✅ Chrome instalado: {ver_chrome}")
    else:
        raise EnvironmentError(
            "Chrome instalado pero no encontrado. "
            "Reiniciá el runtime de Colab y volvé a correr la celda.")
else:
    print(f"  ✅ Chrome OK: {CHROMIUM_BIN}  [{ver_chrome}]")

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.keys import Keys
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────
USERNAME   = "poner minusculas"
PASSWORD   = "password"
BASE_URL   = "https://tourplannx.eurotur.com.ar/TourplanNX_Test"

# Qué comparar — se edita directamente ACÁ, no hace falta exportar ni
# armar ningún CSV/Excel de entrada. Por cada entrada, el script busca
# en Tourplan (Product Search, Location + Supplier + Service Type, sin
# código) TODOS los códigos vigentes de "generico" y de cada supplier
# en "especificos", y les lee la tarifa — ver listar_codigos_supplier.
#
# Cada entrada es una relación genérico/específico independiente — se
# pueden agregar más sin tocar el resto del script, ej. PEAPO (service
# type PJ) o GUIAPO (service type GU) además de TRFPO. "price_code" es
# opcional por entrada (si no está, usa PRICE_CODE_DEFAULT más abajo) —
# no todas las relaciones tienen por qué filtrar por el mismo Price
# Code que TRFPO. "transportistas" sigue aceptándose como alias viejo
# de "especificos" (por compatibilidad con COMPARACIONES ya editados).
COMPARACIONES = [
    {
        "location": "BUE",
        "service_type": "TR",
        "generico": "TRFPO",
        "especificos": ["1TEP01", "6HOUS1"],
        "price_code": "TR",
    },
    # Agregar así cuando haga falta comparar otra relación genérico/
    # específico (sin confirmar todavía contra Tourplan real — ver
    # cabecera del archivo):
    # {
    #     "location": "BUE",
    #     "service_type": "PJ",
    #     "generico": "PEAPO",
    #     "especificos": [...],
    #     "price_code": "TR",  # o "ALL", o lo que corresponda para PEAPO
    # },
    # {
    #     "location": "BUE",
    #     "service_type": "GU",
    #     "generico": "GUIAPO",
    #     "especificos": [...],
    #     "price_code": "TR",
    # },
]
OUTPUT_XLSX = "tarifas_vigentes.xlsx"
OUTPUT_GAP_XLSX = "comparacion_gap.xlsx"
SS_DIR      = "screenshots"
os.makedirs(SS_DIR, exist_ok=True)

# Qué extraer en esta corrida — el genérico de cada relación (TRFPO,
# PEAPO, GUIAPO...) suele cambiar sólo 1-2 veces al año y después queda
# fijo como costo base (ver DISENO.md), así que no hace falta
# re-extraerlo cada vez que se agrega/cambia un específico a comparar:
#   "COMPLETO"          → extrae genérico + específicos (como antes).
#   "SOLO_GENERICO"      → extrae sólo el/los genérico(s) de
#                          COMPARACIONES, los guarda en
#                          tarifas_generico_<GENERICO>_<LOCATION>.xlsx
#                          para reusar después. No calcula comparación
#                          (no hay específico todavía).
#   "SOLO_ESPECIFICOS"   → extrae sólo los "especificos" listados en
#                          COMPARACIONES, y arma la comparación contra
#                          el genérico ya guardado
#                          (tarifas_generico_<GENERICO>_<LOCATION>.xlsx
#                          de una corrida SOLO_GENERICO o COMPLETO
#                          anterior — falla con un error claro si no
#                          existe todavía). "SOLO_TRANSPORTISTA" sigue
#                          aceptándose como alias viejo.
MODO = "COMPLETO"

# Carpeta donde se guarda/lee tarifas_generico_<GENERICO>_<LOCATION>.xlsx
# (el cache de cada genérico). Por default "." (la carpeta donde corre
# el script) — ATENCIÓN si esto corre en Google Colab: el disco de la
# sesión de Colab NO persiste entre sesiones distintas (se pierde al
# desconectarse el runtime). Si vas a correr SOLO_GENERICO un día y
# SOLO_ESPECIFICOS días/semanas después (el caso normal, dado que el
# genérico se mantiene fijo bastante tiempo), poné acá una ruta de
# Google Drive — el script monta Drive solo
# (_montar_drive_si_corresponde), no hace falta un drive.mount(...)
# manual en otra celda:
CACHE_DIR = "."  # ej. "/content/drive/MyDrive/generico-vs-especificos"

# Códigos "genéricos" que no son tarifas reales de ningún servicio —
# siempre cargados en 0 a mano, como placeholder (confirmado por la
# usuaria para 600TRF/700TRF de TRFPO; MINWAT es el mismo caso, visto
# en las muestras). Se descartan en descubrir_cola ANTES de abrir el
# producto — ni siquiera se cuenta el tiempo de Selenium en éstos.
# Editable: agregar el código que haga falta — si PEAPO/GUIAPO u otra
# relación futura tienen sus propios placeholders análogos, van acá
# también (son globales, no por relación, ya que ninguno de estos
# textos/códigos debería aparecer fuera de su contexto de todos modos).
# Mismo criterio que NO_TRANSPORTE en matching_engine.py (Fase 2), pero
# acá no se importa ese módulo para mantener este script autocontenido
# en un solo archivo (Colab).
CODIGOS_EXCLUIR = {"600TRF", "700TRF", "MINWAT"}

# Además del código explícito, se descarta por texto de la descripción
# (cubre variantes que no estén en CODIGOS_EXCLUIR, ej. TRF19/24/42 —
# "TR Generico ...").
PATRONES_EXCLUIR_DESCRIPCION = [
    re.compile(r'MINERAL\s*WATER', re.I),
    re.compile(r'^\s*TRANSFER\s*$', re.I),
    re.compile(r'\bTR\s+GENERICO\b', re.I),
]


def _es_codigo_excluido(codigo, descripcion):
    if (codigo or "").strip().upper() in CODIGOS_EXCLUIR:
        return True
    return any(p.search(descripcion or "") for p in PATRONES_EXCLUIR_DESCRIPCION)

# Rango de fechas a analizar — cada código puede tener varios períodos
# de RATES cargados (ver DISENO.md sección "Períodos"); se procesan
# TODOS los que se superponen con este rango, no sólo uno. Formato
# "dd/Mon/yyyy" (ej. "01/Jan/2026") o "dd/mm/yyyy". Si ambos quedan en
# None, se usa HOY como rango de un solo día — o sea, sólo el período
# vigente en este momento (comportamiento equivalente al de antes de
# que existiera este filtro).
PERIODO_ANALISIS_DESDE = None
PERIODO_ANALISIS_HASTA = None

# Price Code default a filtrar en la lista de RATES antes de leer
# valores, para las relaciones de COMPARACIONES que no traigan su
# propio "price_code" (ver ahí). En modo "All Price Codes" (vista
# agregada default de Tourplan) los valores mostrados NO son los
# persistidos — leer ahí puede devolver 0.0 silenciosamente (hallazgo
# real de tourplan_valorizacion_pkg_v3.py, ver DISENO.md). "ALL"/""/
# None = no filtrar (arriesga el problema anterior si el período tiene
# más de un price code cargado). Confirmado sólo para TRFPO — sin
# confirmar todavía si PEAPO/GUIAPO usan el mismo Price Code "TR" o
# alguno propio (usar "price_code" por entrada en COMPARACIONES si no).
PRICE_CODE_DEFAULT = "TR"

# Límite de códigos a procesar en esta corrida (0 = sin límite, procesa
# toda la cola). Para la primera prueba contra Tourplan real conviene
# un número chico (3-5) — así se valida que los selectores funcionan
# sin quemar tiempo en los ~240 códigos de las muestras si algo falla
# a mitad de camino. Mismo criterio que FASE2_LIMIT en
# tourplan_valorizacion_pkg_v3.py.
LIMIT_PRUEBA = 0

# Tipo de cambio ARS→USD a aplicar sobre las filas cuya moneda leída en
# RATES (columna BUY/SELL CURRENCY) sea ARS — se edita directamente ACÁ,
# no hace falta ningún CSV aparte (ver DISENO.md sección "Moneda"). La
# moneda en sí NO se configura a mano: se lee por línea directamente de
# Tourplan. None = no convertir (las filas en ARS quedan sin TARIFA_USD
# hasta que se complete este valor).
TIPO_CAMBIO_ARS_USD = None

VELOCIDAD = 1.0  # multiplicador de todos los time.sleep — subir si la red es lenta

# Mostrar las capturas inline al correr en Colab/Jupyter (además de
# guardarlas siempre en SS_DIR). Opcional — con muchos códigos ensucia
# mucho el output del notebook, por eso queda en False por default.
MOSTRAR_CAPTURAS = False

_ss_n = [0]
_avisado_sin_ipython = [False]


# ── Helpers base (REUTILIZADO tal cual de los scripts hermanos) ────

def ss(driver, nombre):
    _ss_n[0] += 1
    p = f"{SS_DIR}/{_ss_n[0]:03d}_{nombre[:40]}_{int(time.time())}.png"
    driver.save_screenshot(p)
    print(f"  📸 {os.path.basename(p)}")
    if MOSTRAR_CAPTURAS:
        try:
            from IPython.display import display, Image as IPyImage
            display(IPyImage(p, width=900))
        except ImportError:
            if not _avisado_sin_ipython[0]:
                print("    ⚠ MOSTRAR_CAPTURAS=True pero no hay IPython disponible "
                      "(sólo se ve inline en Colab/Jupyter) — se sigue guardando en disco.")
                _avisado_sin_ipython[0] = True


def dump(driver, nombre):
    p = f"{SS_DIR}/{nombre}_{int(time.time())}.html"
    with open(p, "w", encoding="utf-8") as f:
        f.write(driver.page_source)
    print(f"  💾 HTML: {p}")


def jc(driver, el):
    """JavaScript click — único método confiable en Angular."""
    driver.execute_script("arguments[0].click();", el)


def set_val(driver, el, value):
    driver.execute_script("""
        var inp = arguments[0], val = arguments[1];
        var setter = Object.getOwnPropertyDescriptor(
            window.HTMLInputElement.prototype, 'value').set;
        setter.call(inp, val);
        inp.dispatchEvent(new Event('input',  {bubbles:true}));
        inp.dispatchEvent(new Event('change', {bubbles:true}));
    """, el, value)


def wait(driver, css, t=12):
    return WebDriverWait(driver, t).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, css)))


def crear_driver():
    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--window-size=1366,911")
    opts.add_argument("--disable-gpu")
    if CHROMIUM_BIN:
        opts.binary_location = CHROMIUM_BIN
    drv_path = ChromeDriverManager().install()
    svc = Service(executable_path=drv_path)
    d = webdriver.Chrome(service=svc, options=opts)
    print(f"✅ Driver iniciado (Chrome: {CHROMIUM_BIN or 'default del sistema'})")
    return d


def login(driver):
    print("🔐 Login...")
    driver.get(f"{BASE_URL}/#/login")
    time.sleep(6 * VELOCIDAD)
    ss(driver, "login_page")

    def _campos_visibles():
        return driver.execute_script("""
            function vis(e){return !!(e && (e.offsetWidth||e.offsetHeight
                                      ||e.getClientRects().length)
                                      && !e.disabled);}
            var txt = Array.from(document.querySelectorAll(
                "input[type='text'], input:not([type])")).filter(vis);
            var pwd = Array.from(document.querySelectorAll(
                "input[type='password']")).filter(vis);
            return [txt[0]||null, pwd[0]||null];
        """)

    u_el = p_el = None
    for _ in range(15):
        u_el, p_el = _campos_visibles()
        if u_el and p_el:
            break
        time.sleep(2)
    if not (u_el and p_el):
        ss(driver, "login_sin_campos")
        raise Exception("No aparecieron los campos de login (usuario/password)")

    def _set(el, valor):
        try: el.clear()
        except Exception: pass
        try: el.click()
        except Exception: pass
        try: el.send_keys(valor)
        except Exception:
            set_val(driver, el, valor)

    _set(u_el, USERNAME)
    _set(p_el, PASSWORD)
    time.sleep(0.5)

    clic = driver.execute_script("""
        function vis(e){return !!(e && (e.offsetWidth||e.offsetHeight
                                  ||e.getClientRects().length) && !e.disabled);}
        var b = Array.from(document.querySelectorAll(
            "button.login, button[type='submit'], button")).filter(vis)
            .find(function(x){return /log\\s*in|ingresar|entrar|sign\\s*in/i
                                     .test((x.innerText||'')) ||
                                     x.classList.contains('login');});
        if (b){ b.click(); return (b.innerText||'button.login').trim(); }
        return null;
    """)
    if not clic:
        try: p_el.send_keys(Keys.ENTER)
        except Exception: pass
    time.sleep(8 * VELOCIDAD)
    assert "login" not in driver.current_url.lower(), "Login falló"
    ss(driver, "post_login")
    print("✅ Login OK")


def hamburger(driver):
    img = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, "nav img")))
    jc(driver, img)
    time.sleep(2.5 * VELOCIDAD)


def menu_item(driver, n_or_text):
    if isinstance(n_or_text, int):
        xpath = f"(//nav//ul/li)[{n_or_text}]/div/div"
    else:
        texto = n_or_text.upper()
        xpath = (f"//nav//ul/li[.//*[contains("
                 f"translate(normalize-space(.),'abcdefghijklmnopqrstuvwxyz','ABCDEFGHIJKLMNOPQRSTUVWXYZ'),"
                 f"'{texto}')]]/div/div")
    el = WebDriverWait(driver, 8).until(EC.presence_of_element_located((By.XPATH, xpath)))
    jc(driver, el)
    time.sleep(2 * VELOCIDAD)


# ── Búsqueda de producto (REUTILIZADO de cbd_impact_simulator_pkg.py —
# la versión que abre el resultado y confirma contexto) ────────────

STYPE_SIDEBAR = {
    "HT": "01", "HX": "02", "TF": "03", "EX": "04", "ML": "05",
    "RT": "06", "CR": "07", "FT": "08", "OC": "09", "LN": "10",
    "LP": "11", "MS": "12", "TA": "13", "PR": "14",
}


class ProductoNoEncontrado(Exception):
    pass


def _completar_filtros_busqueda(driver, location, supplier, codigo, service_type):
    """Navega a Product Search y completa Location/Supplier/Code(/Service
    Type) — todo lo que buscar_producto y listar_codigos_supplier
    necesitan en común, antes de divergir en qué hacer con el
    resultado (abrir un producto puntual vs. leer todos los códigos de
    la grilla). Adaptado tal cual de
    cbd_impact_simulator_pkg.py::buscar_producto — ver DISENO.md."""
    st_upper = (service_type or "").strip().upper()

    driver.get(f"{BASE_URL}/#/home")
    time.sleep(2 * VELOCIDAD)
    driver.get(f"{BASE_URL}/#/product")
    time.sleep(5 * VELOCIDAD)

    lupa = wait(driver, "#searchWrapper li:nth-of-type(2) button")
    jc(driver, lupa)
    time.sleep(3 * VELOCIDAD)

    try:
        WebDriverWait(driver, 4).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "div.parameters1 input")))
    except Exception:
        driver.execute_script("""
            var els = Array.from(document.querySelectorAll('li,button,a,div,span'));
            for (var el of els){
                if (!el.offsetParent) continue;
                var t = (el.innerText || '').trim().toUpperCase();
                if (t === 'SELECTION'){ el.click(); return true; }
            }
            return false;
        """)
        time.sleep(2 * VELOCIDAD)

    if st_upper:
        try:
            stype_num = STYPE_SIDEBAR.get(st_upper, "")
            _ok_st = False
            if stype_num:
                _ok_st = driver.execute_script("""
                    var num = arguments[0];
                    for (var li of document.querySelectorAll('li')){
                        if (!li.offsetParent) continue;
                        var t = (li.textContent || '').trim();
                        if (t.indexOf(num) >= 0){ li.click(); return t.slice(0,40); }
                    }
                    return null;
                """, stype_num)
                if _ok_st:
                    time.sleep(1.5 * VELOCIDAD)
            if not _ok_st:
                patron_js = r"(^|[\s\-–—/(])" + re.escape(st_upper) + r"([\s\-–—/)]|$)"
                resultado = None
                for _ in range(3):
                    resultado = driver.execute_script("""
                        var rx = new RegExp(arguments[0]);
                        var elegido = null;
                        for (var li of document.querySelectorAll('li')){
                            if (!li.offsetParent) continue;
                            var t = (li.textContent || '').trim();
                            if (!t || t.length > 80) continue;
                            if (!elegido && rx.test(t.toUpperCase())) elegido = li;
                        }
                        if (elegido){ elegido.click(); return true; }
                        return false;
                    """, patron_js)
                    if resultado:
                        break
                    time.sleep(1.5 * VELOCIDAD)
        except Exception as e:
            print(f"    ⚠ Error seleccionando service type: {e}")

    if location:
        try:
            inp_loc = wait(driver, "div.parameters1 li:nth-of-type(1) input")
            inp_loc.click(); time.sleep(0.3)
            set_val(driver, inp_loc, location)
            time.sleep(1.5 * VELOCIDAD)
            row = WebDriverWait(driver, 5).until(EC.presence_of_element_located(
                (By.CSS_SELECTOR, "tr.selectedRow > td.description")))
            jc(driver, row); time.sleep(1)
        except Exception:
            pass

    if supplier:
        try:
            inp_sup = wait(driver, "div.parameters1 li:nth-of-type(2) input")
            inp_sup.click(); time.sleep(0.3)
            set_val(driver, inp_sup, supplier)
            time.sleep(1.5 * VELOCIDAD)
            row_sup = WebDriverWait(driver, 5).until(EC.presence_of_element_located(
                (By.XPATH, f"//div[contains(@class,'parameters1')]//td[normalize-space(text())='{supplier.upper()}']")))
            jc(driver, row_sup); time.sleep(1)
        except Exception:
            try:
                inp_sup2 = driver.find_element(By.CSS_SELECTOR, "div.parameters1 li:nth-of-type(2) input")
                inp_sup2.send_keys(Keys.TAB); time.sleep(1)
            except Exception:
                pass

    if codigo:
        inp_cod = wait(driver, "div.parameters1 li:nth-of-type(3) input")
        inp_cod.click(); time.sleep(0.3)
        set_val(driver, inp_cod, codigo)
        time.sleep(0.5)

    btn_search = wait(driver, "#productSearchFilter li:nth-of-type(4) button")
    jc(driver, btn_search)
    time.sleep(6 * VELOCIDAD)


def buscar_producto(driver, location, supplier, codigo, service_type=None):
    """Busca Location/Supplier/Code(/ServiceType) en Product Search,
    abre el resultado y confirma que quedó en contexto de producto
    (menú con RATES/UTILITIES/etc.)."""
    codigo = str(codigo).strip() if codigo not in (None, "") else ""
    location = str(location).strip() if location not in (None, "") else ""
    supplier = str(supplier).strip() if supplier not in (None, "") else ""
    st_upper = (service_type or "").strip().upper()
    print(f"\n  📦 Buscando: {location}/{supplier}/{codigo}"
          f"{('/' + st_upper) if st_upper else ''}")

    _completar_filtros_busqueda(driver, location, supplier, codigo, service_type)

    def _en_contexto_producto():
        try:
            img = WebDriverWait(driver, 4).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "nav img")))
            jc(driver, img)
            time.sleep(2 * VELOCIDAD)
            items = driver.execute_script("""
                return Array.from(document.querySelectorAll('nav ul > li')).map(function(li){
                    return (li.querySelector('div div')||li).innerText.trim().split('\\n')[0].toUpperCase();
                });
            """) or []
            ok = "RATES" in items
            jc(driver, img)
            time.sleep(1)
            return ok, items
        except Exception:
            return False, []

    clicked = driver.execute_script("""
        var cod = (arguments[0] || '').toUpperCase();
        var st  = (arguments[1] || '').toUpperCase();
        var rows = Array.from(document.querySelectorAll('table tbody tr'));
        function matchRow(tr){
            var tds = Array.from(tr.querySelectorAll('td'));
            if (cod){
                var hasCod = tds.some(function(td){
                    return td.children.length===0 && td.innerText.trim().toUpperCase()===cod;
                });
                if(!hasCod) return false;
            }
            if(!st) return true;
            return tds.some(function(td){
                return td.children.length===0 && td.innerText.trim().toUpperCase()===st;
            });
        }
        var match = null;
        for (var tr of rows){ if (matchRow(tr)){ match = tr; break; } }
        if (!match && cod){
            for (var tr of rows){
                if ((tr.innerText||'').toUpperCase().includes(cod)){ match = tr; break; }
            }
        }
        if (!match && !cod && rows.length){ match = rows[0]; }
        if (!match) return null;
        var tds = match.querySelectorAll('td');
        var target = tds.length > 4 ? tds[4] : (tds.length > 0 ? tds[tds.length-1] : match);
        target.click();
        return target.innerText.trim().slice(0,50) || 'clicked';
    """, codigo or "", st_upper)

    if not clicked:
        ss(driver, f"ps_sin_resultado_{(codigo or 'sin_codigo')[:10]}")
        raise ProductoNoEncontrado(
            f"Producto no encontrado (location={location!r} supplier={supplier!r} "
            f"codigo={codigo!r} service_type={service_type!r})")

    time.sleep(6 * VELOCIDAD)
    ok_ctx, menu_items = _en_contexto_producto()
    if not ok_ctx:
        ss(driver, f"ps_contexto_incorrecto_{(codigo or 'sin_codigo')[:10]}")
        raise ProductoNoEncontrado(
            f"Click en resultado OK pero no quedó en contexto de producto "
            f"(menú visto: {menu_items}) — codigo={codigo!r}")


def _scrapear_pagina_resultados(driver):
    """Lee la página de resultados de Product Search actualmente
    visible: {items: [{codigo, descripcion}], paginado, total_filas,
    headers}. Columna de código identificada por HEADER de la tabla
    (el `th` cuyo texto es "Code", no cualquier celda corta que
    "parezca" un código — eso confundía la columna Location, ej.
    "BUE", con el código real)."""
    return driver.execute_script(r"""
        var tables = Array.from(document.querySelectorAll('table'));
        var target = null, headers = [];
        for (var t of tables){
            var ths = Array.from(t.querySelectorAll('th')).map(h => h.innerText.trim());
            if (ths.some(h => /CODE/i.test(h))){ target = t; headers = ths; break; }
        }
        if (!target){
            for (var t of tables){
                if (t.querySelectorAll('tbody tr').length){ target = t; break; }
            }
        }
        if (!target) return {items: [], paginado: false, total_filas: 0, headers: []};

        var idxCode = headers.findIndex(h => /^CODE$/i.test(h));
        if (idxCode < 0) idxCode = headers.findIndex(h => /CODE/i.test(h) && !/SERVICE/i.test(h));
        var idxLoc  = headers.findIndex(h => /LOCATION/i.test(h));
        var idxDesc = headers.findIndex(h => /DESCRIPTION/i.test(h));

        var rows = Array.from(target.querySelectorAll('tbody tr'));
        var out = rows.map(function(tr){
            var tds = Array.from(tr.querySelectorAll('td'));
            var codigo, descripcion;
            if (idxCode >= 0 && tds[idxCode]){
                codigo = tds[idxCode].innerText.trim();
            } else {
                var textos = tds.map(function(td){ return td.innerText.trim(); }).filter(Boolean);
                codigo = textos.find(function(t, i){
                    return /^[A-Z0-9]{2,10}$/.test(t) && i !== idxLoc;
                }) || '';
            }
            if (idxDesc >= 0 && tds[idxDesc]){
                descripcion = tds[idxDesc].innerText.trim();
            } else {
                var textos2 = tds.map(function(td){ return td.innerText.trim(); }).filter(Boolean);
                descripcion = textos2.reduce(function(a, b){ return b.length > a.length ? b : a; }, '');
            }
            return {codigo: codigo, descripcion: descripcion};
        }).filter(function(r){ return r.codigo; });

        var paginado = !!document.querySelector("[class*='pagin' i], [class*='pager' i]");
        return {items: out, paginado: paginado, total_filas: rows.length, headers: headers};
    """)


def _hacer_scroll_resultados(driver):
    """Avanza un paso de scroll en el contenedor de resultados —
    confirmado por captura real que la grilla de Product Search NO
    pagina con botón "siguiente": es una lista larga con scroll (virtual
    — sólo las filas visibles están en el DOM, por eso antes se leían
    siempre ~20-22 códigos, el tamaño del viewport, no el catálogo
    completo). Busca el ancestro con scroll real más cercano a la
    tabla y le corre el scrollTop un clientHeight. Devuelve True si
    pudo avanzar, False si ya está en el fondo."""
    return bool(driver.execute_script("""
        function contenedorScroll(){
            var fila = document.querySelector('table tbody tr');
            if (!fila) return null;
            var cur = fila.closest('table');
            while (cur && cur !== document.body){
                if (cur.scrollHeight > cur.clientHeight + 5) return cur;
                cur = cur.parentElement;
            }
            return document.scrollingElement || document.body;
        }
        var c = contenedorScroll();
        if (!c) return false;
        var antes = c.scrollTop;
        c.scrollTop = c.scrollTop + c.clientHeight;
        return c.scrollTop > antes;
    """))


def listar_codigos_supplier(driver, location, supplier, service_type=None):
    """Busca Location/Supplier(/ServiceType) en Product Search SIN
    código (deja ese campo vacío) y devuelve TODOS los resultados —
    scrolleando la grilla (virtual scroll, sin botón de paginación,
    ver _hacer_scroll_resultados) hasta que dejan de aparecer códigos
    nuevos — como lista de {codigo, descripcion}. Es lo que le permite
    a este script no necesitar ningún CSV/Excel de entrada: le alcanza
    con (location, supplier) — ver COMPARACIONES."""
    location = str(location).strip() if location not in (None, "") else ""
    supplier = str(supplier).strip() if supplier not in (None, "") else ""
    print(f"\n  📋 Listando códigos: {location}/{supplier}"
          f"{('/' + service_type.upper()) if service_type else ''}")

    _completar_filtros_busqueda(driver, location, supplier, "", service_type)
    ss(driver, f"listado_{supplier[:15]}")

    items_por_codigo = {}
    headers_vistos = []
    MAX_INTENTOS = 300
    UMBRAL_SIN_NUEVOS = 5
    intentos_sin_nuevos = 0
    for intento in range(1, MAX_INTENTOS + 1):
        resultado = _scrapear_pagina_resultados(driver)
        if not resultado:
            break
        headers_vistos = resultado.get("headers") or headers_vistos
        nuevos = 0
        for item in resultado.get("items", []):
            if item["codigo"] not in items_por_codigo:
                items_por_codigo[item["codigo"]] = item
                nuevos += 1
        if nuevos:
            print(f"    scroll {intento}: {nuevos} códigos nuevos "
                  f"({len(items_por_codigo)} acumulados)")
            intentos_sin_nuevos = 0
        else:
            intentos_sin_nuevos += 1

        # UMBRAL_SIN_NUEVOS scrolls seguidos sin nada nuevo = asumimos
        # que llegamos al final. Subido de 3 a 5 tras un caso real
        # donde la lista quedó corta (HTHT/HTMP/NEZH de TRFPO no se
        # descubrieron en una corrida, sí en otra) — el scroll virtual
        # a veces tarda más en renderizar de lo esperado, sin garantía
        # firme; más margen reduce el riesgo de cortar temprano.
        if intentos_sin_nuevos >= UMBRAL_SIN_NUEVOS:
            break
        if not _hacer_scroll_resultados(driver):
            if intentos_sin_nuevos >= 2:
                break
        time.sleep(1.5 * VELOCIDAD)
    else:
        print(f"    ⚠ Llegué al tope de {MAX_INTENTOS} scrolls para "
              f"{location}/{supplier} — puede haber más códigos sin leer.")

    # Pasada final defensiva: saltar directo al fondo del contenedor
    # (scrollHeight, no +clientHeight) y releer una vez más, por si el
    # loop incremental de arriba cortó antes de llegar al final real.
    driver.execute_script("""
        var fila = document.querySelector('table tbody tr');
        if (!fila) return;
        var cur = fila.closest('table');
        while (cur && cur !== document.body){
            if (cur.scrollHeight > cur.clientHeight + 5){ cur.scrollTop = cur.scrollHeight; return; }
            cur = cur.parentElement;
        }
    """)
    time.sleep(2 * VELOCIDAD)
    resultado_final = _scrapear_pagina_resultados(driver)
    if resultado_final:
        headers_vistos = resultado_final.get("headers") or headers_vistos
        nuevos_final = 0
        for item in resultado_final.get("items", []):
            if item["codigo"] not in items_por_codigo:
                items_por_codigo[item["codigo"]] = item
                nuevos_final += 1
        if nuevos_final:
            print(f"    ⚠ Pasada final (scroll directo al fondo) encontró "
                  f"{nuevos_final} códigos que el loop incremental se había "
                  f"salteado — revisar VELOCIDAD/UMBRAL_SIN_NUEVOS si esto "
                  f"se repite seguido.")

    if not items_por_codigo:
        dump(driver, f"listado_vacio_{supplier[:15]}")
        print(f"    ⚠ No encontré resultados para {location}/{supplier} "
              f"(headers vistos: {headers_vistos})")
        return []
    print(f"    → {len(items_por_codigo)} códigos encontrados en total "
          f"(headers: {headers_vistos})")
    return list(items_por_codigo.values())


# ── Períodos de RATES (REUTILIZADO tal cual de
# tourplan_valorizacion_pkg_v3.py::actualizar_rates_servicio_madre —
# _leer_periodos/_parse_rate_period/_seleccionar_price_code, selectores
# confirmados contra Tourplan real: td.tpcol-rateperiod,
# td.tpcol-pricecodecode, #priceCodeModeSelected, #priceCode) ────────

MESES_ES = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
            "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}


def parsear_fecha(txt):
    """'01/Apr/2026', '01/04/2026' o '2026-04-01' → datetime. None si
    no matchea ningún formato conocido."""
    if not txt:
        return None
    if isinstance(txt, datetime):
        return txt
    txt = str(txt).strip()
    if len(txt) >= 10 and txt[4] == "-" and txt[7] == "-":
        try:
            return datetime.fromisoformat(txt[:10])
        except Exception:
            pass
    p = txt.split("/")
    if len(p) == 3:
        try:
            dia = int(p[0])
            mes_raw = p[1]
            mes = MESES_ES.get(mes_raw[:3].capitalize()) or int(mes_raw)
            anio_raw = int(p[2])
            anio = anio_raw + 2000 if anio_raw < 100 else anio_raw
            return datetime(anio, mes, dia)
        except Exception:
            pass
    return None


def _parse_rate_period(texto):
    """Parsea '01/Apr/2026 - 31/Aug/2026' → (datetime, datetime)."""
    m = re.search(r'(\d+/\w+/\d+)\s*[-–]\s*(\d+/\w+/\d+)', texto or "")
    if m:
        return parsear_fecha(m.group(1)), parsear_fecha(m.group(2))
    return None, None


def _seleccionar_price_code(driver, code, etiqueta=""):
    """En la lista de RATES (o dentro de un período ya abierto), pasa
    el modo a 'Selected Price Code' y elige el price code indicado.
    En modo 'All Price Codes' (default) la grilla es la vista agregada
    y lo que se lee ahí NO es lo persistido (puede leerse 0.0 sin
    ningún error — hallazgo real del script hermano). code vacío o
    "ALL"/"UNASSIGNED" no hace nada (se deja la vista default)."""
    code = (code or "").strip().upper()
    if not code or code in ("ALL", "TODOS", "*", "UNASSIGNED"):
        return False
    try:
        driver.execute_script("""
            var lbl = document.querySelector('label[for="priceCodeModeSelected"]');
            var inp = document.getElementById('priceCodeModeSelected');
            if (lbl) lbl.click();
            if (inp) inp.click();
        """)
        time.sleep(1.5 * VELOCIDAD)
        driver.execute_script("""
            var dd = document.getElementById('priceCode');
            if (!dd) return;
            var inp = dd.querySelector('input[type="text"], input');
            if (inp){ inp.focus(); inp.click(); } else { dd.click(); }
        """)
        time.sleep(1.0 * VELOCIDAD)
        driver.execute_script("""
            var dd = document.getElementById('priceCode');
            if (!dd) return;
            var inp = dd.querySelector('input[type="text"], input');
            if (!inp) return;
            var setter = Object.getOwnPropertyDescriptor(
                window.HTMLInputElement.prototype, 'value').set;
            setter.call(inp, arguments[0]);
            inp.dispatchEvent(new Event('input', {bubbles:true}));
            inp.dispatchEvent(new Event('keyup', {bubbles:true}));
        """, code)
        time.sleep(1.2 * VELOCIDAD)
        elegido = driver.execute_script("""
            var pc = arguments[0].toUpperCase();
            var dd = document.getElementById('priceCode');
            if (!dd) return null;
            function vis(e){return !!(e.offsetWidth||e.offsetHeight||e.getClientRects().length);}
            function match(tr){
                var cod = tr.querySelector('td.code');
                var des = tr.querySelector('td.description');
                var ct = cod ? (cod.innerText||cod.textContent||'').trim().toUpperCase() : '';
                var dt = des ? (des.innerText||des.textContent||'').trim().toUpperCase() : '';
                if (ct===pc || dt===pc || dt.indexOf(pc+' ')===0 ||
                    dt.indexOf(pc+'-')===0 || dt.indexOf(pc+' -')===0)
                    return {tr:tr, txt:ct + ' | ' + dt};
                return null;
            }
            var rows = Array.from(dd.querySelectorAll('table tr'));
            var cand = null;
            for (var tr of rows){ if (vis(tr)){ var m=match(tr); if(m){cand=m;break;} } }
            if (!cand){ for (var tr2 of rows){ var m2=match(tr2); if(m2){cand=m2;break;} } }
            if (!cand) return null;
            try{ cand.tr.scrollIntoView({block:'center'}); }catch(e){}
            var des = cand.tr.querySelector('td.description');
            var cod = cand.tr.querySelector('td.code');
            (des||cod||cand.tr).click();
            cand.tr.click();
            return cand.txt;
        """, code)
        time.sleep(2.5 * VELOCIDAD)
        print(f"    Price Code '{code}' seleccionado ({etiqueta}): {elegido}")
        if not elegido:
            ss(driver, f"pricecode_fail_{(etiqueta or code)[:12]}")
        return bool(elegido)
    except Exception as e:
        print(f"    ⚠ No pude seleccionar Price Code '{code}' ({etiqueta}): {e}")
        return False


def _leer_periodos_rates(driver):
    """Lista de {texto, desde, hasta, price_code, moneda} de la grilla
    de RATES actualmente visible (lista de períodos, no el detalle de
    uno). Cada fila trae la fecha en td.tpcol-rateperiod y su price
    code en td.tpcol-pricecodecode (selectores reales confirmados). La
    moneda (Buy/Sell Currency, iguales por default — confirmado por la
    usuaria) se lee ACÁ, en esta lista, y no de la grilla de un período
    ya abierto: ahí la moneda viene embebida en el texto del header
    (ej. "USD\\nGROUP COST"), que no es una fuente confiable en
    general. Se busca la columna por header ("CURRENCY", prioridad a
    "BUY") en la misma tabla que contiene los td.tpcol-rateperiod."""
    filas = driver.execute_script("""
        var celda0 = document.querySelector('td.tpcol-rateperiod');
        var tabla = celda0 ? celda0.closest('table') : null;
        var headers = tabla ? Array.from(tabla.querySelectorAll('th')).map(function(h){
            return h.innerText.trim();
        }) : [];
        var idxCcy = headers.findIndex(function(h){ return /BUY/i.test(h) && /CURR/i.test(h); });
        if (idxCcy < 0) idxCcy = headers.findIndex(function(h){ return /CURR/i.test(h); });

        var out = [];
        document.querySelectorAll('td.tpcol-rateperiod').forEach(function(d){
            var tr = d.closest('tr');
            var p = tr ? tr.querySelector('td.tpcol-pricecodecode') : null;
            var tds = tr ? Array.from(tr.querySelectorAll('td')) : [];
            var moneda = (idxCcy >= 0 && tds[idxCcy]) ? tds[idxCcy].innerText.trim() : '';
            out.push({texto: (d.innerText||'').trim(),
                      price_code: p ? (p.innerText||'').trim() : '',
                      moneda: moneda, headers: headers});
        });
        return out;
    """) or []
    out = []
    for f in filas:
        desde, hasta = _parse_rate_period(f["texto"])
        out.append({"texto": f["texto"], "price_code": f["price_code"],
                    "moneda": f["moneda"].strip().upper(), "desde": desde, "hasta": hasta})
    if filas and not any(f["moneda"] for f in filas):
        print(f"    ⚠ No encontré columna de moneda (BUY/SELL CURRENCY) en la lista "
              f"de períodos — headers vistos: {filas[0].get('headers')}")
    return out


def _periodos_en_rango(periodos, desde_str, hasta_str):
    """Índices de `periodos` cuyo rango de fechas se superpone con
    [desde_str, hasta_str]. Si ambos son None, usa HOY como rango de
    un solo día (o sea, el período vigente ahora mismo). Períodos sin
    fecha parseable (texto inesperado) se descartan con aviso, no se
    incluyen "por si acaso"."""
    objetivo_desde = parsear_fecha(desde_str) if desde_str else datetime.now()
    objetivo_hasta = parsear_fecha(hasta_str) if hasta_str else datetime.now()
    idxs = []
    for i, p in enumerate(periodos):
        if not (p["desde"] and p["hasta"]):
            print(f"    ⚠ Período con fecha no parseable, se descarta: {p['texto']!r}")
            continue
        if p["desde"] <= objetivo_hasta and p["hasta"] >= objetivo_desde:
            idxs.append(i)
    return idxs


# ── Lectura de la grilla RATES del componente (REUTILIZADO de
# tourplan_valorizacion_pkg_v3.py — _leer_tabla_rates/_extraer_valores_ad,
# generalizado para devolver el rango de pax desde/hasta en vez de sólo
# la etiqueta de fila) ──────────────────────────────────────────────

def _leer_tabla_rates(driver):
    return driver.execute_script("""
        for(var t of document.querySelectorAll('table')){
            var ths = Array.from(t.querySelectorAll('th'));
            if(ths.some(h => h.innerText.includes('GROUP COST') ||
                              h.innerText.includes('COST'))){
                return {
                    headers: ths.map(h => h.innerText.trim()),
                    rows: Array.from(t.querySelectorAll('tbody tr')).map(tr => ({
                        celdas: Array.from(tr.querySelectorAll('td')).map(td => td.innerText.trim()),
                        inputs: Array.from(tr.querySelectorAll('input')).map(i => i.value.trim())
                    }))
                };
            }
        }
        return null;
    """)


def _listar_tablas_pagina(driver):
    """Diagnóstico: headers de TODAS las tablas visibles en la página
    actual — para cuando _leer_tabla_rates no encuentra ninguna con
    columna COST, y hay que ver a ciegas qué hay en pantalla (ej. una
    lista de períodos en vez de la grilla de tarifas) sin necesitar el
    HTML dump completo."""
    return driver.execute_script("""
        return Array.from(document.querySelectorAll('table')).map(function(t){
            return Array.from(t.querySelectorAll('th')).map(function(h){ return h.innerText.trim(); });
        });
    """)


def _detectar_moneda_headers(headers):
    """Fallback si la lista de períodos no trajo moneda: la grilla de
    un período abierto la muestra embebida en el texto del header (ej.
    "USD\\nGROUP COST") — no es la fuente principal (ver
    _leer_periodos_rates) porque no se confirmó que sea confiable en
    todos los casos, pero sirve de respaldo."""
    for h in headers:
        m = re.match(r'^([A-Z]{3})\b', h.strip())
        if m:
            return m.group(1)
    return ""


def _parsear_numero(raw):
    """Convierte el texto de una celda de tarifa a float. Tourplan usa
    formato US/UK para TODAS las monedas que se vieron hasta ahora —
    coma = separador de miles, punto = decimal — confirmado con un
    valor real de 6HOUS1 en ARS: "85,000" (ochenta y cinco mil pesos;
    si la coma fuera decimal sería "85 pesos", sin sentido para un
    transfer). Antes se intentaba adivinar cuál separador era el
    decimal por cuál aparecía más a la derecha (asumiendo que ARS podía
    venir en formato argentino, punto de miles + coma decimal) — con
    "85,000" (sin punto) ese heurístico lo interpretaba mal como 85.0.
    Ahora se asume siempre esta única convención, para USD y para ARS."""
    s = str(raw).replace("$", "").replace(" ", "").strip()
    if not s:
        return None
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _extraer_filas_ad(tabla, codigo, moneda_periodo=""):
    """De una tabla ya leída con _leer_tabla_rates, devuelve las filas
    de rango de pax adulto (AD) como {pax_desde, pax_hasta, tarifa,
    moneda} — igual criterio que _extraer_valores_ad del script
    hermano. La moneda se recibe del período (ver
    _leer_periodos_rates) en vez de buscarse en esta grilla; sólo si
    viene vacía se intenta el fallback de header. No toca
    período/price code, eso lo maneja el llamador."""
    headers = tabla["headers"]
    idx_svc = next((i for i, h in enumerate(headers) if "SERVICE" in h.upper()), 0)
    idx_cost = next((i for i, h in enumerate(headers)
                      if "GROUP COST" in h.upper() and "FIT" not in h.upper()), 1)
    pat_rango = re.compile(r'(\d+)\s*[-–]\s*(\d+)')

    moneda = moneda_periodo or _detectar_moneda_headers(headers)
    if not moneda:
        print(f"    ⚠ No pude determinar la moneda de {codigo} (ni en la lista de "
              f"períodos ni en los headers de la grilla) — headers: {headers}")

    out = []
    for row in tabla["rows"]:
        celdas = row["celdas"]
        if not celdas:
            continue
        svc = celdas[idx_svc] if idx_svc < len(celdas) else ""
        if not svc.strip():
            continue
        m = pat_rango.search(svc)
        if not (m and re.search(r'\bAD\b', svc.upper())):
            continue
        pax_hasta = int(m.group(2))
        if pax_hasta == 9999:
            # Catch-all de pax abierto hacia arriba (ej. "42-9999 AD")
            # — confirmado por la usuaria que siempre da 0, no aporta
            # nada a la comparación y sólo infla la cantidad de filas.
            continue
        raw = row["inputs"][0] if (row["inputs"] and row["inputs"][0]) else (
            celdas[idx_cost] if idx_cost < len(celdas) else "")
        valor = _parsear_numero(raw)
        out.append({
            "pax_desde": int(m.group(1)),
            "pax_hasta": pax_hasta,
            "tarifa": round(valor, 2) if valor is not None else None,
            "moneda": moneda,
        })
    return out


def _abrir_lista_rates(driver, codigo, etiqueta, price_code):
    """Navega (o vuelve a navegar) a RATES desde el producto en
    contexto, y aplica el filtro de price_code si corresponde. Se llama
    de nuevo por cada período a abrir porque, al entrar al detalle de
    un período, no queda forma confirmada de "volver" a la lista salvo
    renavegar (mismo patrón de re-navegación que usa
    tourplan_valorizacion_pkg_v3.py en su verificación post-SAVE)."""
    hamburger(driver)
    menu_item(driver, "RATES")
    time.sleep(3 * VELOCIDAD)
    if price_code:
        _seleccionar_price_code(driver, price_code, etiqueta=etiqueta)
        time.sleep(1.5 * VELOCIDAD)
    ss(driver, f"rates_lista_{codigo[:10]}")


def leer_tarifa_vigente_componente(driver, codigo, price_code=None):
    """Abre el tab RATES del producto ya en contexto (ver
    buscar_producto), identifica qué período(s) caen dentro de
    [PERIODO_ANALISIS_DESDE, PERIODO_ANALISIS_HASTA] (por defecto, sólo
    el período vigente hoy — ver esa constante), abre cada uno y lee su
    grilla de tarifas. Devuelve una lista de {pax_desde, pax_hasta,
    tarifa, moneda, periodo_desde, periodo_hasta, price_code} — un
    conjunto de filas AD por cada período que matcheó.

    price_code: filtro de Price Code a aplicar en RATES — viene de la
    relación genérico/específico de COMPARACIONES (cada una puede tener
    el suyo, ej. "TR" para TRFPO; otra relación como PEAPO/GUIAPO puede
    necesitar uno distinto o "ALL"). None/"" = usa PRICE_CODE_DEFAULT.

    Moneda: cada período de Rates tiene columnas BUY CURRENCY y SELL
    CURRENCY, por defecto cargadas iguales — confirmado por la usuaria,
    así que alcanza con leer una (se prioriza BUY si ambas existen). No
    hay fallback a config manual: si esta columna no aparece en la
    grilla real, la fila queda con moneda vacía y se avisa (ver
    DISENO.md), en vez de asumir una moneda por proveedor."""
    price_code = price_code if price_code is not None else PRICE_CODE_DEFAULT
    _abrir_lista_rates(driver, codigo, etiqueta=f"lista {codigo}", price_code=price_code)

    periodos = _leer_periodos_rates(driver)
    if not periodos:
        dump(driver, f"rates_sin_periodos_{codigo[:10]}")
        tablas_vistas = _listar_tablas_pagina(driver)
        print(f"    ⚠ No encontré lista de períodos de RATES para {codigo}")
        print(f"      Tablas visibles en la página (headers): {tablas_vistas}")
        return []

    idxs = _periodos_en_rango(periodos, PERIODO_ANALISIS_DESDE, PERIODO_ANALISIS_HASTA)
    if not idxs:
        print(f"    ⚠ Ningún período de {codigo} cae dentro del rango configurado "
              f"(PERIODO_ANALISIS_DESDE/HASTA) — períodos vistos: "
              f"{[p['texto'] for p in periodos]}")
        return []
    print(f"    → {len(idxs)}/{len(periodos)} período(s) dentro del rango: "
          f"{[periodos[i]['texto'] for i in idxs]}")

    out = []
    for n, idx in enumerate(idxs):
        if n > 0:
            # Renavegar: el DOM de la lista se perdió al entrar al período anterior.
            _abrir_lista_rates(driver, codigo, etiqueta=f"lista {codigo} (período {n+1})",
                                price_code=price_code)
        periodo = periodos[idx]
        filas_periodo = driver.find_elements(By.CSS_SELECTOR, "td.tpcol-rateperiod")
        if idx >= len(filas_periodo):
            print(f"    ⚠ El período {periodo['texto']!r} de {codigo} ya no está en la "
                  f"posición esperada tras renavegar — se salta.")
            continue
        jc(driver, filas_periodo[idx])
        time.sleep(5 * VELOCIDAD)
        ss(driver, f"rates_periodo_{codigo[:10]}_{n+1}")
        if price_code:
            _seleccionar_price_code(driver, price_code, etiqueta=f"período {codigo}")

        tabla = _leer_tabla_rates(driver)
        for _ in range(3):
            if tabla:
                break
            time.sleep(2 * VELOCIDAD)
            tabla = _leer_tabla_rates(driver)
        if not tabla:
            dump(driver, f"rates_sin_tabla_{codigo[:10]}_{n+1}")
            tablas_vistas = _listar_tablas_pagina(driver)
            print(f"    ⚠ No encontré tabla de RATES (con columna COST) para {codigo} "
                  f"período {periodo['texto']!r}")
            print(f"      Tablas visibles en la página (headers): {tablas_vistas}")
            continue

        for fila in _extraer_filas_ad(tabla, codigo, moneda_periodo=periodo.get("moneda", "")):
            fila["periodo_desde"] = periodo["desde"].strftime("%d/%m/%Y") if periodo["desde"] else ""
            fila["periodo_hasta"] = periodo["hasta"].strftime("%d/%m/%Y") if periodo["hasta"] else ""
            fila["price_code"] = periodo["price_code"]
            out.append(fila)
    return out


# ── Cola de trabajo y moneda ─────────────────────────────────────────

def descubrir_cola(driver, comparaciones, incluir_generico=True, incluir_especificos=True):
    """A partir de COMPARACIONES (una entrada por relación genérico/
    específico — location + supplier genérico + lista de específicos,
    ver COMPARACIONES), busca en Tourplan los códigos vigentes de cada
    supplier (listar_codigos_supplier) y devuelve la cola de trabajo
    completa — un (LOCATION, SUPPLIER, CODIGO, SERVICE_TYPE, ES_GENERICO,
    PRICE_CODE) por código encontrado. Requiere sesión ya logueada.

    incluir_generico/incluir_especificos permiten limitar qué se extrae
    esta corrida (ver MODO) — para no re-extraer el genérico completo
    cuando sólo hace falta actualizar un específico."""
    cola = []
    for comp in comparaciones:
        location = comp["location"]
        service_type = comp.get("service_type", "")
        price_code = comp.get("price_code", PRICE_CODE_DEFAULT)
        especificos = comp.get("especificos", comp.get("transportistas", []))
        suppliers = []
        if incluir_generico:
            suppliers.append((comp["generico"], True))
        if incluir_especificos:
            suppliers += [(s, False) for s in especificos]
        for supplier, es_generico in suppliers:
            items = listar_codigos_supplier(driver, location, supplier, service_type)
            for item in items:
                if _es_codigo_excluido(item["codigo"], item.get("descripcion", "")):
                    print(f"    (excluido, no es un servicio real: {item['codigo']} "
                          f"— {item.get('descripcion', '')!r})")
                    continue
                cola.append({
                    "location": location,
                    "supplier": supplier,
                    "codigo": item["codigo"],
                    "descripcion": item.get("descripcion", ""),
                    "service_type": service_type,
                    "es_generico": es_generico,
                    "generico": comp["generico"],
                    "price_code": price_code,
                })
    return cola


COLS_TARIFAS = ["SUPPLIER", "PRODUCT_CODE", "ES_GENERICO", "GENERICO", "PERIODO_DESDE",
                "PERIODO_HASTA", "PRICE_CODE", "PAX_DESDE", "PAX_HASTA", "TARIFA_VIGENTE",
                "MONEDA", "TARIFA_USD", "LOCATION", "TIMESTAMP"]
# GENERICO: a qué supplier genérico corresponde esta relación (para
# una fila ES_GENERICO=True, es su propio SUPPLIER; para una fila
# específica, el "generico" de la entrada de COMPARACIONES de la que
# vino) — necesario para no mezclar relaciones distintas que comparten
# location (ej. TRFPO y PEAPO, ambos en BUE) al armar la comparación.


def _montar_drive_si_corresponde():
    """Si CACHE_DIR apunta a Google Drive (ej. "/content/drive/...") y
    esto corre en Colab, monta Drive automáticamente — no hace falta
    un `drive.mount(...)` manual en otra celda. `drive.mount` ya es
    idempotente (si ya está montado no hace nada raro), así que se
    puede llamar en cada corrida sin problema. La primera vez en un
    navegador/cuenta nueva, Colab de todos modos va a pedir un click de
    autorización (flujo de seguridad de Google, no de este script) —
    eso no se puede saltear."""
    if not str(CACHE_DIR).startswith("/content/drive"):
        return
    try:
        from google.colab import drive
    except ImportError:
        print(f"    ⚠ CACHE_DIR={CACHE_DIR!r} apunta a Google Drive, pero esto "
              f"no está corriendo en Colab (no hay módulo google.colab) — "
              f"dejalo montado a mano o usá una ruta local.")
        return
    print("📂 Montando Google Drive...")
    drive.mount('/content/drive')


def _archivo_cache_generico(generico, location):
    return os.path.join(CACHE_DIR, f"tarifas_generico_{generico}_{location}.xlsx")


def guardar_cache_generico(filas_salida):
    """Guarda, por (supplier genérico, location), las filas
    ES_GENERICO=True de esta corrida en
    tarifas_generico_<GENERICO>_<LOCATION>.xlsx — para que una corrida
    futura en MODO=SOLO_ESPECIFICOS no tenga que re-extraer el genérico
    completo (ver DISENO.md, "Resuelto: desacoplar la extracción de
    TRFPO..."). Un archivo distinto por supplier (TRFPO/PEAPO/GUIAPO...)
    para que no se pisen entre sí en la misma location. Se llama
    siempre que esta corrida haya extraído genérico (COMPLETO o
    SOLO_GENERICO), pisando el cache anterior de ese supplier+location."""
    if CACHE_DIR not in (".", ""):
        os.makedirs(CACHE_DIR, exist_ok=True)
    por_generico_location = {}
    for f in filas_salida:
        if f["ES_GENERICO"]:
            por_generico_location.setdefault((f["SUPPLIER"], f["LOCATION"]), []).append(f)
    for (generico, location), filas in por_generico_location.items():
        wb = Workbook()
        hoja = wb.active
        hoja.title = "TARIFAS_GENERICO"
        hoja.append(COLS_TARIFAS)
        for c in hoja[1]:
            c.font = Font(bold=True)
        for fila in filas:
            hoja.append([fila.get(c, "") for c in COLS_TARIFAS])
        for i, c in enumerate(COLS_TARIFAS, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)
        archivo = _archivo_cache_generico(generico, location)
        wb.save(archivo)
        print(f"💾 Cache de {generico} guardado: {archivo} ({len(filas)} filas)")


def cargar_cache_generico(comparaciones):
    """Carga tarifas_generico_<GENERICO>_<LOCATION>.xlsx por cada
    relación (supplier genérico, location) de COMPARACIONES — usado en
    MODO=SOLO_ESPECIFICOS para no re-extraer el genérico. Falla con un
    mensaje claro si falta el archivo de alguna relación (hace falta
    haber corrido MODO=SOLO_GENERICO o COMPLETO al menos una vez antes,
    para esa relación)."""
    filas = []
    relaciones = {(c["generico"], c["location"]) for c in comparaciones}
    for generico, location in relaciones:
        archivo = _archivo_cache_generico(generico, location)
        if not os.path.exists(archivo):
            raise FileNotFoundError(
                f"No encontré {archivo}. Con MODO=SOLO_ESPECIFICOS hace falta "
                f"haber corrido antes MODO=SOLO_GENERICO (o COMPLETO) al menos "
                f"una vez para {generico!r} en {location!r}, así queda guardado "
                f"el cache que esta corrida necesita. Si esto corre en Colab y "
                f"CACHE_DIR='.' (default), el archivo se pierde entre sesiones "
                f"distintas del runtime — montá Google Drive y apuntá CACHE_DIR "
                f"ahí para que el cache sobreviva de una sesión a otra (ver "
                f"comentario junto a CACHE_DIR al principio del script).")
        wb = load_workbook(archivo, data_only=True)
        filas_hoja = list(wb.active.iter_rows(values_only=True))
        header = filas_hoja[0]
        for row in filas_hoja[1:]:
            filas.append(dict(zip(header, row)))
        print(f"📂 Cache de {generico} cargado: {archivo} ({len(filas_hoja) - 1} filas)")
    return filas


def convertir_a_usd(tarifa, moneda, tipo_cambio):
    """None si no se puede convertir con confianza: falta la tarifa,
    la moneda vino vacía (columna CURRENCY no encontrada), la moneda no
    es ARS ni USD, o es ARS pero no hay tipo de cambio cargado — mejor
    dejar el dato en blanco para revisión manual que inventar un
    valor."""
    if tarifa is None or not moneda:
        return None
    if moneda == "USD":
        return round(tarifa, 2)
    if moneda == "ARS":
        return round(tarifa / tipo_cambio, 2) if tipo_cambio else None
    return None


# ── Fase 3 embebida: comparación de gap (versión en Python puro de
# comparacion_gap.py — sin pandas, para que este script siga siendo
# autocontenido en un solo archivo/una sola corrida de Colab. Misma
# lógica que ese archivo standalone; si se cambia una, replicar en la
# otra. Ver DISENO.md) ──────────────────────────────────────────────

COLS_GAP = [
    "LOCATION", "CODIGO_GENERICO", "SUPPLIER_ESPECIFICO", "CODIGO_ESPECIFICO",
    "PAX_DESDE", "PAX_HASTA",
    "PERIODO_ESPECIFICO_DESDE", "PERIODO_ESPECIFICO_HASTA",
    "PERIODO_GENERICO_DESDE", "PERIODO_GENERICO_HASTA",
    "TARIFA_GENERICO_USD", "TARIFA_ESPECIFICO_USD",
    "DIFERENCIA_USD", "DIFERENCIA_PCT", "FLAGS",
]


def _mejor_prefijo_generico(codigo, codigos_generico_desc):
    """Longest-prefix-match código específico → código genérico. Mismo
    algoritmo que matching_engine.py::mejor_prefijo_trfpo (Fase 2) —
    nunca dependió de nada específico de TRFPO (es prefijo de string
    puro), así que sirve igual para cualquier otra relación genérico/
    específico (PEAPO, GUIAPO, la que sea) sin cambios."""
    candidatos = [c for c in codigos_generico_desc if codigo.startswith(c)]
    if not candidatos:
        return None, []
    return max(candidatos, key=len), candidatos


def _dias_superposicion(desde1, hasta1, desde2, hasta2):
    """Copiado de comparacion_gap.py::_dias_superposicion. Días de
    superposición entre dos rangos de fechas — 0 (no negativo) si no
    se superponen o si falta alguna fecha, para poder usar max() sin
    excepciones por None."""
    if not (desde1 and hasta1 and desde2 and hasta2):
        return 0
    inicio = max(desde1, desde2)
    fin = min(hasta1, hasta2)
    return max(0, (fin - inicio).days)


def construir_comparacion_gap(filas):
    """A partir de filas_salida (la misma lista de dicts que se
    escribe en tarifas_vigentes.xlsx, ya en memoria — no hace falta
    releer el Excel) calcula el gap genérico vs. cada código
    específico. Cada fila específica sólo se compara contra los
    genéricos de SU MISMA relación (columna `GENERICO`, ver
    COLS_TARIFAS) — no contra cualquier genérico que comparta su
    location — para no mezclar entre sí relaciones distintas que
    convivan en la misma location (ej. TRFPO y PEAPO ambos en BUE). Si
    a una fila específica le falta `GENERICO` (dato viejo, de antes de
    que existiera esta columna), cae a compararse contra TODOS los
    genéricos de su location, como se hacía antes. Ver comparacion_gap.py
    para el detalle de banderas y criterio de match (idéntico acá, sólo
    sin pandas)."""
    por_location = {}
    for f in filas:
        por_location.setdefault(f["LOCATION"], []).append(f)

    salida = []
    for location, filas_loc in por_location.items():
        filas_generico_todos = [f for f in filas_loc if f["ES_GENERICO"]]
        filas_especifico = [f for f in filas_loc if not f["ES_GENERICO"]]

        for fe in filas_especifico:
            generico_asociado = fe.get("GENERICO")
            filas_generico = (
                [f for f in filas_generico_todos if f["SUPPLIER"] == generico_asociado]
                if generico_asociado else filas_generico_todos
            )
            codigos_generico = sorted({f["PRODUCT_CODE"] for f in filas_generico}, key=len, reverse=True)

            base = {
                "LOCATION": location,
                "SUPPLIER_ESPECIFICO": fe["SUPPLIER"],
                "CODIGO_ESPECIFICO": fe["PRODUCT_CODE"],
                "PAX_DESDE": fe["PAX_DESDE"], "PAX_HASTA": fe["PAX_HASTA"],
                "PERIODO_ESPECIFICO_DESDE": fe["PERIODO_DESDE"],
                "PERIODO_ESPECIFICO_HASTA": fe["PERIODO_HASTA"],
                "PERIODO_GENERICO_DESDE": "", "PERIODO_GENERICO_HASTA": "",
                "TARIFA_ESPECIFICO_USD": fe["TARIFA_USD"],
            }
            mejor, candidatos = _mejor_prefijo_generico(fe["PRODUCT_CODE"], codigos_generico)
            flags = []
            if len(candidatos) > 1:
                flags.append("COLISION_REVISAR")
            if mejor is None:
                salida.append({**base, "CODIGO_GENERICO": "", "TARIFA_GENERICO_USD": None,
                               "DIFERENCIA_USD": None, "DIFERENCIA_PCT": None,
                               "FLAGS": ",".join(flags + ["SIN_MATCH_GENERICO"])})
                continue

            candidatas = [
                f for f in filas_generico
                if f["PRODUCT_CODE"] == mejor
                and f["PAX_DESDE"] is not None and f["PAX_HASTA"] is not None
                and fe["PAX_DESDE"] is not None and fe["PAX_HASTA"] is not None
                and f["PAX_DESDE"] <= fe["PAX_HASTA"] and f["PAX_HASTA"] >= fe["PAX_DESDE"]
            ]
            if not candidatas:
                salida.append({**base, "CODIGO_GENERICO": mejor, "TARIFA_GENERICO_USD": None,
                               "DIFERENCIA_USD": None, "DIFERENCIA_PCT": None,
                               "FLAGS": ",".join(flags + ["SIN_GENERICO_PARA_ESE_PAX"])})
                continue

            # Más de un período del genérico cargado para este mismo
            # rango de pax (ej. cambió su tarifa a mitad del período
            # del específico): se usa el de MAYOR superposición de
            # fechas con el período del específico, y se marca con
            # MULTIPLES_PERIODOS_GENERICO para no descartar el resto en
            # silencio — el/la que revise sabe que hay otro(s)
            # período(s) del genérico en tarifas_vigentes.xlsx para ese
            # mismo código/pax.
            if len(candidatas) > 1:
                flags.append("MULTIPLES_PERIODOS_GENERICO")

            def _dias(f):
                return _dias_superposicion(
                    parsear_fecha(fe["PERIODO_DESDE"]), parsear_fecha(fe["PERIODO_HASTA"]),
                    parsear_fecha(f["PERIODO_DESDE"]), parsear_fecha(f["PERIODO_HASTA"]))

            fila_generico = max(candidatas, key=_dias)
            if _dias(fila_generico) == 0:
                flags.append("SIN_SUPERPOSICION_DE_PERIODO")

            tarifa_generico = fila_generico["TARIFA_USD"]
            tarifa_especifico = fe["TARIFA_USD"]
            diff_usd = diff_pct = None
            if tarifa_generico is not None and tarifa_especifico is not None:
                diff_usd = round(tarifa_especifico - tarifa_generico, 2)
                if tarifa_generico:
                    diff_pct = round((tarifa_especifico / tarifa_generico - 1) * 100, 2)
            else:
                flags.append("SIN_TARIFA_USD_PARA_COMPARAR")

            salida.append({**base, "CODIGO_GENERICO": mejor,
                           "PERIODO_GENERICO_DESDE": fila_generico["PERIODO_DESDE"],
                           "PERIODO_GENERICO_HASTA": fila_generico["PERIODO_HASTA"],
                           "TARIFA_GENERICO_USD": tarifa_generico,
                           "DIFERENCIA_USD": diff_usd, "DIFERENCIA_PCT": diff_pct,
                           "FLAGS": ",".join(flags)})
    return salida


# ── Main ──────────────────────────────────────────────────────────

def main():
    modo = (MODO or "COMPLETO").strip().upper()
    modo = {"SOLO_TRANSPORTISTA": "SOLO_ESPECIFICOS"}.get(modo, modo)  # alias viejo
    if modo not in ("SOLO_GENERICO", "SOLO_ESPECIFICOS", "COMPLETO"):
        raise ValueError(f"MODO inválido: {MODO!r} — usar SOLO_GENERICO, "
                          f"SOLO_ESPECIFICOS o COMPLETO")
    incluir_generico = modo in ("SOLO_GENERICO", "COMPLETO")
    incluir_especificos = modo in ("SOLO_ESPECIFICOS", "COMPLETO")
    print(f"MODO={modo} (genérico: {'sí' if incluir_generico else 'no'}, "
          f"específicos: {'sí' if incluir_especificos else 'no'})")

    _montar_drive_si_corresponde()

    tipo_cambio = TIPO_CAMBIO_ARS_USD
    if tipo_cambio is None:
        print(f"⚠ Sin TIPO_CAMBIO_ARS_USD cargado (constante al principio del "
              f"script) — las filas en ARS quedarán sin TARIFA_USD hasta completarlo.")

    driver = crear_driver()
    filas_salida = []
    try:
        login(driver)

        cola = descubrir_cola(driver, COMPARACIONES, incluir_generico, incluir_especificos)
        if LIMIT_PRUEBA:
            print(f"\n⚠ LIMIT_PRUEBA={LIMIT_PRUEBA} — procesando sólo los primeros "
                  f"{LIMIT_PRUEBA} códigos de {len(cola)} descubiertos. Poner "
                  f"LIMIT_PRUEBA=0 para correr la cola completa.")
            cola = cola[:LIMIT_PRUEBA]
        print(f"Cola de trabajo: {len(cola)} códigos a leer.")

        for i, item in enumerate(cola):
            print(f"\n[{i+1}/{len(cola)}] {item['supplier']}/{item['codigo']}")
            try:
                buscar_producto(driver, item["location"], item["supplier"],
                                 item["codigo"], service_type=item["service_type"])
                tarifas = leer_tarifa_vigente_componente(driver, item["codigo"],
                                                          price_code=item["price_code"])
                if not tarifas:
                    print(f"    ⚠ Sin tarifas leídas para {item['codigo']}")
                for t in tarifas:
                    if t["moneda"] not in ("ARS", "USD"):
                        print(f"    ⚠ Moneda inesperada {t['moneda']!r} en "
                              f"{item['codigo']} (pax {t['pax_desde']}-{t['pax_hasta']})"
                              f" — no se convierte a USD")
                    filas_salida.append({
                        "SUPPLIER": item["supplier"],
                        "PRODUCT_CODE": item["codigo"],
                        "ES_GENERICO": item["es_generico"],
                        "GENERICO": item["generico"],
                        "PERIODO_DESDE": t.get("periodo_desde", ""),
                        "PERIODO_HASTA": t.get("periodo_hasta", ""),
                        "PRICE_CODE": t.get("price_code", ""),
                        "PAX_DESDE": t["pax_desde"],
                        "PAX_HASTA": t["pax_hasta"],
                        "TARIFA_VIGENTE": t["tarifa"],
                        "MONEDA": t["moneda"],
                        "TARIFA_USD": convertir_a_usd(t["tarifa"], t["moneda"], tipo_cambio),
                        "LOCATION": item["location"],
                        "TIMESTAMP": datetime.now().isoformat(timespec="seconds"),
                    })
            except ProductoNoEncontrado as e:
                print(f"    ❌ {e}")
            except Exception as e:
                print(f"    ❌ Error inesperado en {item['codigo']}: {e}")
                ss(driver, f"error_{item['codigo'][:10]}")
    finally:
        driver.quit()

    wb = Workbook()
    hoja = wb.active
    hoja.title = "TARIFAS_VIGENTES"
    hoja.append(COLS_TARIFAS)
    for c in hoja[1]:
        c.font = Font(bold=True)
    for fila in filas_salida:
        hoja.append([fila.get(c, "") for c in COLS_TARIFAS])
    for i, c in enumerate(COLS_TARIFAS, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)
    wb.save(OUTPUT_XLSX)
    print(f"\n✅ {len(filas_salida)} filas escritas en {OUTPUT_XLSX}")

    if incluir_generico:
        guardar_cache_generico(filas_salida)

    if modo == "SOLO_GENERICO":
        print("\nMODO=SOLO_GENERICO: no se calcula comparación de gap en esta "
              "corrida (no se extrajo ningún específico). Corré "
              "MODO=SOLO_ESPECIFICOS para comparar contra alguno usando este cache.")
        return

    filas_para_comparacion = filas_salida
    if modo == "SOLO_ESPECIFICOS":
        filas_para_comparacion = cargar_cache_generico(COMPARACIONES) + filas_salida

    print("\n🔄 Calculando comparación de gap (Fase 3)...")
    comparacion = construir_comparacion_gap(filas_para_comparacion)
    wb_gap = Workbook()
    hoja_gap = wb_gap.active
    hoja_gap.title = "COMPARACION_GAP"
    hoja_gap.append(COLS_GAP)
    for c in hoja_gap[1]:
        c.font = Font(bold=True)
    idx_pct = COLS_GAP.index("DIFERENCIA_PCT") + 1
    for fila in comparacion:
        hoja_gap.append([fila.get(c, "") for c in COLS_GAP])
    for fila_excel in hoja_gap.iter_rows(min_row=2, min_col=idx_pct, max_col=idx_pct):
        for celda in fila_excel:
            if isinstance(celda.value, (int, float)):
                celda.number_format = '0.00"%"'
    for i, c in enumerate(COLS_GAP, start=1):
        hoja_gap.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 2)
    wb_gap.save(OUTPUT_GAP_XLSX)
    print(f"✅ {len(comparacion)} filas escritas en {OUTPUT_GAP_XLSX}")

    con_flags = [f for f in comparacion if f["FLAGS"]]
    if con_flags:
        conteo = {}
        for f in con_flags:
            conteo[f["FLAGS"]] = conteo.get(f["FLAGS"], 0) + 1
        print(f"\n{len(con_flags)} filas con alguna bandera (revisar):")
        for flag, n in sorted(conteo.items(), key=lambda x: -x[1]):
            print(f"    {flag}: {n}")


if __name__ == "__main__":
    main()
